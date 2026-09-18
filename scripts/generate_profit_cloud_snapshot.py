"""Generate a public, browser-readable profit snapshot from TikTok exports.

The order export is reduced to fields needed by the calculation engine. Original
order IDs and creator handles are deliberately excluded from the published file.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl


SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PUBLIC_VERSION = "2026-09-profit-cloud-1"


def normalized_sku(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"\s+", "", text).casefold()


def clean_text(value: object) -> str:
    return str(value or "").strip()


def clean_number(value: object) -> float | None:
    if value is None or value == "":
        return None
    text = str(value).strip().replace(",", "").replace("฿", "").replace("%", "")
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def parse_order_date(value: object) -> str:
    text = clean_text(value)
    for date_format in (
        "%d/%m/%Y %H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%m/%d/%Y %H:%M:%S",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, date_format).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def _cell_value(cell: ET.Element) -> str:
    value = cell.find(f"{SHEET_NS}v")
    if value is not None:
        return value.text or ""
    inline = cell.find(f"{SHEET_NS}is/{SHEET_NS}t")
    return inline.text if inline is not None else ""


def read_sparse_first_sheet(path: Path) -> tuple[dict[str, str], list[dict[str, str]]]:
    """Read exports whose worksheet dimension incorrectly reports only A1."""
    with zipfile.ZipFile(path) as archive:
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        relation_xml = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relations = {node.attrib["Id"]: node.attrib["Target"] for node in relation_xml}
        first_sheet = workbook_xml.find(f"{SHEET_NS}sheets/{SHEET_NS}sheet")
        if first_sheet is None:
            raise ValueError(f"{path.name}: workbook has no worksheet")
        relation_id = first_sheet.attrib[f"{REL_NS}id"]
        target = relations[relation_id].lstrip("/")
        if not target.startswith("xl/"):
            target = f"xl/{target}"
        worksheet_bytes = archive.read(target)

    cells_by_row: dict[int, dict[str, str]] = {}
    for _, cell in ET.iterparse(io.BytesIO(worksheet_bytes), events=("end",)):
        if cell.tag != f"{SHEET_NS}c":
            continue
        coordinate = cell.attrib.get("r", "")
        match = re.fullmatch(r"([A-Z]+)(\d+)", coordinate)
        if match:
            column, row_number = match.group(1), int(match.group(2))
            cells_by_row.setdefault(row_number, {})[column] = _cell_value(cell)
        cell.clear()

    headers = cells_by_row.get(1, {})
    column_by_header = {value: column for column, value in headers.items()}
    rows = []
    for row_number in sorted(cells_by_row):
        if row_number <= 2:
            continue
        row_cells = cells_by_row[row_number]
        rows.append({header: row_cells.get(column, "") for header, column in column_by_header.items()})
    return headers, rows


def seller_revenue(row: dict[str, str]) -> float | None:
    before = clean_number(row.get("SKU Subtotal Before Discount"))
    seller_discount = clean_number(row.get("SKU Seller Discount"))
    if before is not None and seller_discount is not None:
        return max(0.0, before - seller_discount)
    return clean_number(row.get("SKU Subtotal After Discount"))


def build_public_orders(path: Path, store: str) -> list[dict[str, object]]:
    _, source_rows = read_sparse_first_sheet(path)
    order_aliases: dict[str, str] = {}
    records: list[dict[str, object]] = []
    for row in source_rows:
        raw_order_id = clean_text(row.get("Order ID"))
        if not re.fullmatch(r"\d{6,}", raw_order_id):
            continue
        if raw_order_id not in order_aliases:
            order_aliases[raw_order_id] = f"PUB-{len(order_aliases) + 1:06d}"
        quantity = clean_number(row.get("Quantity"))
        record = {
            "store": store,
            "orderId": order_aliases[raw_order_id],
            "productId": "",
            "status": clean_text(row.get("Order Status")),
            "skuId": clean_text(row.get("SKU ID")),
            "sellerSku": clean_text(row.get("Seller SKU")),
            "productName": clean_text(row.get("Product Name")),
            "qty": quantity if quantity is not None else 1,
            "returnQty": clean_number(row.get("Sku Quantity of return")) or 0,
            "unitPrice": clean_number(row.get("SKU Unit Original Price")),
            "subtotalBeforeDiscount": clean_number(row.get("SKU Subtotal Before Discount")),
            "platformDiscount": clean_number(row.get("SKU Platform Discount")),
            "sellerDiscount": clean_number(row.get("SKU Seller Discount")),
            "subtotalAfterDiscount": clean_number(row.get("SKU Subtotal After Discount")),
            "dealPrice": seller_revenue(row),
            "orderAmount": clean_number(row.get("Order Amount")),
            "refund": clean_number(row.get("Order Refund Amount")),
            "date": parse_order_date(row.get("Created Time")),
            "weightKg": clean_number(row.get("Weight(kg)")),
        }
        records.append(record)
    if not records:
        raise ValueError(f"{path.name}: no order rows were parsed")
    return records


def _workbook_skus(path: Path) -> tuple[list[dict[str, object]], dict[str, object]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    detail = workbook["产品定价利润明细"]
    rows = detail.iter_rows(min_row=1, values_only=True)
    first_rows = [next(rows) for _ in range(3)]
    settings = first_rows[1]
    flags = {
        "miaosha": clean_text(settings[2]) == "是",
        "live": clean_text(settings[5]) == "是",
        "affAd": clean_text(settings[7]) == "是",
    }
    skus = []
    for row in rows:
        sku = clean_text(row[0])
        if not sku:
            continue
        skus.append({
            "sku": sku,
            "base": sku.split()[0],
            "weightKg": clean_number(row[1]),
            "cost": clean_number(row[2]),
            "activityPrice": clean_number(row[3]),
            "suggestedRetailPrice": clean_number(row[4]),
        })

    rates = {
        "transaction": 0.0321,
        "shopCommission": 0.107,
        "affCommission": 0.10,
        "affAdCommission": 0.03,
        "miaosha": 0.0321,
        "live": 0.0321,
        "growth": 0.0803,
        "infra": 0.0015,
        "staff": 0.06,
    }
    rate_labels = {
        "交易手续费": "transaction",
        "Shop 佣金": "shopCommission",
        "Shop佣金": "shopCommission",
        "联盟店铺广告佣金": "affAdCommission",
        "联盟广告佣金": "affAdCommission",
        "联盟佣金": "affCommission",
        "秒杀": "miaosha",
        "直播": "live",
        "电商增长": "growth",
        "基础设施": "infra",
        "人员": "staff",
        "综合成本": "staff",
    }
    for row in workbook["费率参数表"].iter_rows(values_only=True):
        name = clean_text(row[0])
        value = clean_number(row[1])
        if not name or value is None:
            continue
        for label, key in rate_labels.items():
            if label in name:
                rates[key] = value / 100 if value > 1 else value
                break

    tiers = []
    for row in workbook["运费阶梯表"].iter_rows(min_row=3, values_only=True):
        lower = clean_number(row[0])
        upper = clean_number(row[1])
        if lower is None and upper is None:
            continue
        tiers.append({"lo": lower or 0, "hi": upper, "net": clean_number(row[4]) or 0})
    return skus, {"flags": flags, "rates": rates, "tiers": tiers}


def build_public_pricing(cost_map_path: Path, pricing_path: Path, generated_at: str) -> dict[str, object]:
    detailed_skus, workbook_data = _workbook_skus(pricing_path)
    cost_workbook = openpyxl.load_workbook(cost_map_path, read_only=True, data_only=True)
    cost_sheet = cost_workbook.active
    cost_by_sku: dict[str, tuple[float, str]] = {}
    for sku, cost, *_ in cost_sheet.iter_rows(min_row=2, values_only=True):
        key = normalized_sku(sku)
        numeric_cost = clean_number(cost)
        if key and numeric_cost is not None:
            cost_by_sku[key] = (numeric_cost, clean_text(sku))

    merged = []
    detailed_keys = set()
    for sku in detailed_skus:
        key = normalized_sku(sku["sku"])
        detailed_keys.add(key)
        output = dict(sku)
        if key in cost_by_sku:
            output["cost"] = cost_by_sku[key][0]
            output["costSource"] = "SKU 成本映射表"
        merged.append(output)
    for key, (cost, source_sku) in cost_by_sku.items():
        if key in detailed_keys:
            continue
        merged.append({
            "sku": source_sku,
            "base": source_sku.split()[0],
            "weightKg": None,
            "cost": cost,
            "activityPrice": None,
            "suggestedRetailPrice": None,
            "costSource": "SKU 成本映射表",
        })

    return {
        "importedAt": generated_at,
        "fileName": "INSPIRE PURIFY 成本与定价公开快照",
        "sourceType": "published-snapshot",
        "sourceFiles": ["SKU 成本映射表", "TikTok 泰国站产品价格利润核算表"],
        "costMapFiles": ["SKU 成本映射表"],
        "publishedSnapshot": True,
        "hasShippingTiers": bool(workbook_data["tiers"]),
        "hasRateSheet": True,
        "hasActivityPrices": any(sku["activityPrice"] is not None for sku in merged),
        "flags": workbook_data["flags"],
        "rates": workbook_data["rates"],
        "tiers": workbook_data["tiers"],
        "skus": merged,
    }


def update_cloud_manifest(path: Path, generated_at: str) -> None:
    prefix = "window.TIKTOK_CLOUD_SNAPSHOT="
    suffix = ";window.REAL_STORE_DATA=window.TIKTOK_CLOUD_SNAPSHOT.realStoreData;"
    source = path.read_text(encoding="utf-8-sig").strip()
    if not source.startswith(prefix) or not source.endswith(suffix):
        raise ValueError(f"{path.name}: unsupported cloud snapshot wrapper")
    snapshot = json.loads(source[len(prefix):-len(suffix)])
    snapshot["version"] = PUBLIC_VERSION
    snapshot["generatedAt"] = generated_at
    snapshot.setdefault("v33Urls", {})["orders"] = f"./data/cloud-orders.json?v={PUBLIC_VERSION}"
    snapshot["pricingUrl"] = f"./data/cloud-pricing.json?v={PUBLIC_VERSION}"
    serialized = json.dumps(snapshot, ensure_ascii=False, separators=(",", ":"))
    path.write_text(
        f"window.TIKTOK_CLOUD_SNAPSHOT={serialized};window.REAL_STORE_DATA=window.TIKTOK_CLOUD_SNAPSHOT.realStoreData;",
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--orders", type=Path, required=True)
    parser.add_argument("--cost-map", type=Path, required=True)
    parser.add_argument("--pricing", type=Path, required=True)
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument("--store", default="INSPIRE PURIFY")
    args = parser.parse_args()

    generated_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    data_directory = args.repo / "data"
    orders = build_public_orders(args.orders, args.store)
    pricing = build_public_pricing(args.cost_map, args.pricing, generated_at)
    (data_directory / "cloud-orders.json").write_text(
        json.dumps(orders, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    (data_directory / "cloud-pricing.json").write_text(
        json.dumps(pricing, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    update_cloud_manifest(data_directory / "cloud-import.js", generated_at)
    print(json.dumps({
        "orders": len(orders),
        "publicOrders": len({row["orderId"] for row in orders}),
        "pricingSkus": len(pricing["skus"]),
        "dateMin": min(row["date"] for row in orders if row["date"]),
        "dateMax": max(row["date"] for row in orders if row["date"]),
        "version": PUBLIC_VERSION,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
